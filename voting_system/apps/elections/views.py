from __future__ import annotations

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Border, Font, Side
from rest_framework.views import APIView

from voting_system.apps.accounts.models import MembershipRole, OrgMembership
from voting_system.apps.analytics.services import election_summary, election_turnout_timeseries
from voting_system.apps.audit.services import log_event
from voting_system.apps.ballots.services import get_results_payload
from voting_system.apps.common.responses import error_response, success_response
from voting_system.apps.elections.models import Candidate, Election, ElectionStatus, Post, ResultsVisibility
from voting_system.apps.elections.serializers import (
    CandidateSerializer,
    ElectionCreateSerializer,
    ElectionDuplicateSerializer,
    ElectionSerializer,
    ElectionStatusSerializer,
    PostSerializer,
)
from voting_system.apps.elections.services import ElectionTransitionError, change_election_status, duplicate_election
from voting_system.apps.organizations.models import Organization


def _resolve_org(request):
    org_id = request.headers.get("X-Org-Id") or request.query_params.get("org_id")
    if not org_id:
        return None
    return Organization.objects.filter(id=org_id).first()


def _membership_for_request(request, organization):
    if getattr(request.user, "is_system_admin", False):
        return None
    return OrgMembership.objects.filter(user=request.user, organization=organization, is_active=True).first()


def _can_view_org(request, organization):
    membership = _membership_for_request(request, organization)
    if getattr(request.user, "is_system_admin", False):
        return True, membership
    if not membership:
        return False, None
    return membership.role in (
        MembershipRole.ORG_ADMIN,
        MembershipRole.ELECTION_MANAGER,
        MembershipRole.RESULTS_VIEWER,
    ), membership


def _can_edit_org(request, organization):
    membership = _membership_for_request(request, organization)
    if getattr(request.user, "is_system_admin", False):
        return True, membership
    if not membership:
        return False, None
    return membership.role in (MembershipRole.ORG_ADMIN, MembershipRole.ELECTION_MANAGER), membership


def _election_write_payload(data):
    allowed_fields = set(ElectionCreateSerializer.Meta.fields)
    return {key: value for key, value in data.items() if key in allowed_fields}


def _ensure_election_not_archived(election):
    if election.status == ElectionStatus.ARCHIVED:
        return error_response("Archived elections are read-only", status_code=400)
    return None


def _build_excel(election, results):
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    summary = election_summary(election)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    summary_rows = [
        ("Metric", "Value"),
        ("Election Title", election.title),
        ("Organization", election.organization.name),
        ("Opens At", str(election.opens_at or "")),
        ("Closes At", str(election.closes_at or "")),
        ("Tokens Generated", summary["tokens_generated"]),
        ("Tokens Used", summary["tokens_used"]),
        ("Turnout %", summary["turnout_percentage"]),
        ("Ballots Submitted", summary["ballots_submitted"]),
    ]
    for row_idx, (label, value) in enumerate(summary_rows, start=1):
        ws_summary.cell(row=row_idx, column=1, value=label)
        ws_summary.cell(row=row_idx, column=2, value=value)
        ws_summary.cell(row=row_idx, column=1).border = border
        ws_summary.cell(row=row_idx, column=2).border = border
        if row_idx == 1:
            ws_summary.cell(row=row_idx, column=1).font = Font(bold=True)
            ws_summary.cell(row=row_idx, column=2).font = Font(bold=True)
    ws_summary.freeze_panes = "A2"

    for post in results["posts"]:
        ws = wb.create_sheet((post["post_title"] or "Post")[:31])
        headers = ["Candidate", "Votes", "Percentage"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.border = border
        ws.freeze_panes = "A2"

        for row_index, candidate in enumerate(post["candidates"], start=2):
            ws.cell(row=row_index, column=1, value=candidate["name"]).border = border
            ws.cell(row=row_index, column=2, value=candidate["votes"]).border = border
            ws.cell(row=row_index, column=3, value=round(candidate["percentage"], 2)).border = border

        for col_idx in range(1, 4):
            width = 12
            for row_idx in range(1, ws.max_row + 1):
                cell_value = ws.cell(row=row_idx, column=col_idx).value
                width = max(width, len(str(cell_value or "")) + 2)
            ws.column_dimensions[chr(64 + col_idx)].width = min(width, 60)

    for col_idx in range(1, 3):
        width = 12
        for row_idx in range(1, ws_summary.max_row + 1):
            cell_value = ws_summary.cell(row=row_idx, column=col_idx).value
            width = max(width, len(str(cell_value or "")) + 2)
        ws_summary.column_dimensions[chr(64 + col_idx)].width = min(width, 64)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


class OrgElectionListCreateView(APIView):
    def get(self, request):
        organization = _resolve_org(request)
        if not organization:
            return error_response("Organization context missing", status_code=400)
        allowed, _ = _can_view_org(request, organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        elections = Election.objects.filter(organization=organization)
        return success_response(ElectionSerializer(elections, many=True).data)

    def post(self, request):
        organization = _resolve_org(request)
        if not organization:
            return error_response("Organization context missing", status_code=400)
        allowed, _ = _can_edit_org(request, organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        serializer = ElectionCreateSerializer(data=_election_write_payload(request.data))
        serializer.is_valid(raise_exception=True)
        fallback_auto_logout_seconds = serializer.validated_data.get(
            "voter_auto_logout_seconds",
            organization.voter_auto_logout_seconds,
        )
        fallback_voter_results_after_vote_enabled = serializer.validated_data.get(
            "voter_results_after_vote_enabled",
            False,
        )
        election = serializer.save(
            organization=organization,
            status="DRAFT",
            created_by=request.user,
            voter_auto_logout_seconds=fallback_auto_logout_seconds,
            voter_results_after_vote_enabled=fallback_voter_results_after_vote_enabled,
        )
        log_event(
            actor=request.user,
            organization=organization,
            action="ELECTION_CREATED",
            target_type="election",
            target_id=str(election.id),
            metadata={"title": election.title, "slug": election.slug},
        )
        return success_response(ElectionSerializer(election).data, "Election created", 201)


class OrgElectionDetailView(APIView):
    def get(self, request, election_id):
        election = Election.objects.select_related("organization").filter(id=election_id).first()
        if not election:
            return error_response("Election not found", status_code=404)
        allowed, _ = _can_view_org(request, election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        data = ElectionSerializer(election).data
        data["summary"] = election_summary(election)
        data["turnout_series"] = election_turnout_timeseries(election)
        return success_response(data)

    def patch(self, request, election_id):
        election = Election.objects.select_related("organization").filter(id=election_id).first()
        if not election:
            return error_response("Election not found", status_code=404)
        archived_error = _ensure_election_not_archived(election)
        if archived_error:
            return archived_error
        allowed, _ = _can_edit_org(request, election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        serializer = ElectionCreateSerializer(election, data=_election_write_payload(request.data), partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        log_event(
            actor=request.user,
            organization=election.organization,
            action="ELECTION_UPDATED",
            target_type="election",
            target_id=str(election.id),
            metadata={"title": election.title},
        )
        return success_response(ElectionSerializer(election).data, "Election updated")

    def delete(self, request, election_id):
        election = Election.objects.select_related("organization").filter(id=election_id).first()
        if not election:
            return error_response("Election not found", status_code=404)
        archived_error = _ensure_election_not_archived(election)
        if archived_error:
            return archived_error

        membership = _membership_for_request(request, election.organization)
        if not getattr(request.user, "is_system_admin", False):
            if not membership or membership.role != MembershipRole.ORG_ADMIN:
                return error_response("Forbidden", status_code=403)

        if election.status in (ElectionStatus.SCHEDULED, ElectionStatus.LIVE):
            return error_response("Only non-active elections can be deleted", status_code=400)

        organization = election.organization
        target_id = str(election.id)
        metadata = {"title": election.title, "slug": election.slug, "status": election.status}
        election.delete()
        log_event(
            actor=request.user,
            organization=organization,
            action="ELECTION_DELETED",
            target_type="election",
            target_id=target_id,
            metadata=metadata,
        )
        return success_response(message="Election deleted")


class OrgElectionStatusActionView(APIView):
    def post(self, request, election_id):
        election = Election.objects.select_related("organization").filter(id=election_id).first()
        if not election:
            return error_response("Election not found", status_code=404)
        archived_error = _ensure_election_not_archived(election)
        if archived_error:
            return archived_error
        can_edit, membership = _can_edit_org(request, election.organization)
        if not can_edit:
            return error_response("Forbidden", status_code=403)

        serializer = ElectionStatusSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        actor_role = membership.role if membership else MembershipRole.ORG_ADMIN
        try:
            election = change_election_status(
                election=election,
                action=serializer.validated_data["action"],
                actor=request.user,
                actor_role=actor_role,
            )
        except ElectionTransitionError as exc:
            return error_response(str(exc), status_code=400)
        return success_response(ElectionSerializer(election).data, "Election status updated")


class OrgPostListCreateView(APIView):
    def get(self, request, election_id):
        election = Election.objects.select_related("organization").filter(id=election_id).first()
        if not election:
            return error_response("Election not found", status_code=404)
        allowed, _ = _can_view_org(request, election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        return success_response(PostSerializer(Post.objects.filter(election=election), many=True).data)

    def post(self, request, election_id):
        election = Election.objects.select_related("organization").filter(id=election_id).first()
        if not election:
            return error_response("Election not found", status_code=404)
        archived_error = _ensure_election_not_archived(election)
        if archived_error:
            return archived_error
        allowed, _ = _can_edit_org(request, election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        serializer = PostSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        post = serializer.save(election=election)
        log_event(
            actor=request.user,
            organization=election.organization,
            action="POST_CREATED",
            target_type="post",
            target_id=str(post.id),
            metadata={"title": post.title},
        )
        return success_response(PostSerializer(post).data, "Post created", 201)


class OrgPostDetailView(APIView):
    def patch(self, request, post_id):
        post = Post.objects.select_related("election__organization").filter(id=post_id).first()
        if not post:
            return error_response("Post not found", status_code=404)
        archived_error = _ensure_election_not_archived(post.election)
        if archived_error:
            return archived_error
        allowed, _ = _can_edit_org(request, post.election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        serializer = PostSerializer(post, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        log_event(
            actor=request.user,
            organization=post.election.organization,
            action="POST_UPDATED",
            target_type="post",
            target_id=str(post.id),
            metadata={"title": post.title},
        )
        return success_response(PostSerializer(post).data, "Post updated")

    def delete(self, request, post_id):
        post = Post.objects.select_related("election__organization").filter(id=post_id).first()
        if not post:
            return error_response("Post not found", status_code=404)
        archived_error = _ensure_election_not_archived(post.election)
        if archived_error:
            return archived_error
        allowed, _ = _can_edit_org(request, post.election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        org = post.election.organization
        target_id = str(post.id)
        post.delete()
        log_event(
            actor=request.user,
            organization=org,
            action="POST_DELETED",
            target_type="post",
            target_id=target_id,
            metadata={},
        )
        return success_response(message="Post deleted")


class OrgCandidateListCreateView(APIView):
    def get(self, request, post_id):
        post = Post.objects.select_related("election__organization").filter(id=post_id).first()
        if not post:
            return error_response("Post not found", status_code=404)
        allowed, _ = _can_view_org(request, post.election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        return success_response(CandidateSerializer(Candidate.objects.filter(post=post), many=True).data)

    def post(self, request, post_id):
        post = Post.objects.select_related("election__organization").filter(id=post_id).first()
        if not post:
            return error_response("Post not found", status_code=404)
        archived_error = _ensure_election_not_archived(post.election)
        if archived_error:
            return archived_error
        allowed, _ = _can_edit_org(request, post.election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        serializer = CandidateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        candidate = serializer.save(post=post)
        log_event(
            actor=request.user,
            organization=post.election.organization,
            action="CANDIDATE_CREATED",
            target_type="candidate",
            target_id=str(candidate.id),
            metadata={"name": candidate.name},
        )
        return success_response(CandidateSerializer(candidate).data, "Candidate created", 201)


class OrgCandidateDetailView(APIView):
    def patch(self, request, candidate_id):
        candidate = Candidate.objects.select_related("post__election__organization").filter(id=candidate_id).first()
        if not candidate:
            return error_response("Candidate not found", status_code=404)
        archived_error = _ensure_election_not_archived(candidate.post.election)
        if archived_error:
            return archived_error
        allowed, _ = _can_edit_org(request, candidate.post.election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        serializer = CandidateSerializer(candidate, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        log_event(
            actor=request.user,
            organization=candidate.post.election.organization,
            action="CANDIDATE_UPDATED",
            target_type="candidate",
            target_id=str(candidate.id),
            metadata={"name": candidate.name},
        )
        return success_response(CandidateSerializer(candidate).data, "Candidate updated")

    def delete(self, request, candidate_id):
        candidate = Candidate.objects.select_related("post__election__organization").filter(id=candidate_id).first()
        if not candidate:
            return error_response("Candidate not found", status_code=404)
        archived_error = _ensure_election_not_archived(candidate.post.election)
        if archived_error:
            return archived_error
        allowed, _ = _can_edit_org(request, candidate.post.election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        org = candidate.post.election.organization
        target_id = str(candidate.id)
        candidate.delete()
        log_event(
            actor=request.user,
            organization=org,
            action="CANDIDATE_DELETED",
            target_type="candidate",
            target_id=target_id,
            metadata={},
        )
        return success_response(message="Candidate deleted")


class OrgElectionResultsView(APIView):
    def get(self, request, election_id):
        election = Election.objects.select_related("organization").prefetch_related("posts__candidates").filter(id=election_id).first()
        if not election:
            return error_response("Election not found", status_code=404)
        allowed, _ = _can_view_org(request, election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        if election.results_visibility == ResultsVisibility.HIDDEN_UNTIL_CLOSED and election.status not in ("CLOSED", "ARCHIVED"):
            return error_response("Results are hidden until election closes", status_code=403)
        if election.results_visibility == ResultsVisibility.LIVE_ALLOWED and election.status not in ("LIVE", "CLOSED", "ARCHIVED"):
            return error_response("Results are unavailable for this election status", status_code=403)
        return success_response(get_results_payload(election))


class OrgElectionPublishResultsView(APIView):
    def post(self, request, election_id):
        election = Election.objects.select_related("organization").filter(id=election_id).first()
        if not election:
            return error_response("Election not found", status_code=404)
        archived_error = _ensure_election_not_archived(election)
        if archived_error:
            return archived_error
        allowed, _ = _can_edit_org(request, election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        if election.status not in ("CLOSED", "ARCHIVED"):
            return error_response("Results can only be published after the election is closed", status_code=400)
        publish = bool(request.data.get("publish", True))
        election.publish_results = publish
        election.save(update_fields=["publish_results", "updated_at"])
        log_event(
            actor=request.user,
            organization=election.organization,
            action="RESULTS_PUBLISHED" if publish else "RESULTS_UNPUBLISHED",
            target_type="election",
            target_id=str(election.id),
            metadata={"publish_results": publish},
        )
        return success_response(ElectionSerializer(election).data, "Results publish status updated")


class OrgElectionExportExcelView(APIView):
    def get(self, request, election_id):
        election = Election.objects.select_related("organization").prefetch_related("posts__candidates").filter(id=election_id).first()
        if not election:
            return error_response("Election not found", status_code=404)
        allowed, _ = _can_view_org(request, election.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)
        if election.results_visibility == ResultsVisibility.HIDDEN_UNTIL_CLOSED and election.status not in ("CLOSED", "ARCHIVED"):
            return error_response("Results are hidden until election closes", status_code=403)
        if election.results_visibility == ResultsVisibility.LIVE_ALLOWED and election.status not in ("LIVE", "CLOSED", "ARCHIVED"):
            return error_response("Results are unavailable for this election status", status_code=403)
        results = get_results_payload(election)
        payload = _build_excel(election, results)
        log_event(
            actor=request.user,
            organization=election.organization,
            action="RESULTS_EXPORTED",
            target_type="election",
            target_id=str(election.id),
            metadata={"format": "excel"},
        )
        response = HttpResponse(payload, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{election.slug}_results.xlsx"'
        return response


class OrgElectionDuplicateView(APIView):
    def post(self, request, election_id):
        source = Election.objects.select_related("organization").filter(id=election_id).first()
        if not source:
            return error_response("Election not found", status_code=404)
        allowed, _ = _can_edit_org(request, source.organization)
        if not allowed:
            return error_response("Forbidden", status_code=403)

        serializer = ElectionDuplicateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        duplicate = duplicate_election(
            source=source,
            actor=request.user,
            title=serializer.validated_data.get("title"),
            slug=serializer.validated_data.get("slug"),
        )
        return success_response(ElectionSerializer(duplicate).data, "Election duplicated", 201)

# Create your views here.
